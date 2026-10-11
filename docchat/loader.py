"""Read PDF / DOCX / TXT / MD files and cut them into searchable chunks.

Everything works from raw bytes in memory, so there are no temp files,
no path problems and no file-name quirks on Windows.
"""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from typing import Callable, List, Optional, Tuple

ALLOWED_TYPES = ["pdf", "docx", "txt", "md"]


@dataclass
class Chunk:
    text: str
    source: str
    page: Optional[int] = None
    kind: str = "text"  # "text" or "table"


Block = Tuple[Optional[int], str, str]  # (page, kind, text)


# --------------------------------------------------------------------------
# Readers
# --------------------------------------------------------------------------
def _read_pdf(data: bytes, extract_tables: bool) -> List[Block]:
    import pymupdf  # PyMuPDF

    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        if doc.needs_pass:
            raise ValueError("the PDF is password-protected")
        blocks: List[Block] = []
        for number, page in enumerate(doc, start=1):
            text = page.get_text("text") or ""
            if text.strip():
                blocks.append((number, "text", text))
            if extract_tables:
                try:
                    for table in page.find_tables().tables:
                        md = (table.to_markdown() or "").strip()
                        if md:
                            blocks.append((number, "table", md))
                except Exception:
                    pass  # table detection is a bonus; never fail the upload
        return blocks
    finally:
        doc.close()


def _read_docx(data: bytes) -> List[Block]:
    import docx

    document = docx.Document(io.BytesIO(data))
    blocks: List[Block] = []
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    if paragraphs:
        blocks.append((None, "text", "\n\n".join(paragraphs)))
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            ncols = len(table.rows[0].cells)
            rows.insert(1, "|" + " --- |" * ncols)
            blocks.append((None, "table", "\n".join(rows)))
    return blocks


def _read_text(data: bytes) -> List[Block]:
    text = data.decode("utf-8-sig", errors="replace")
    return [(None, "text", text)] if text.strip() else []


def read_file(name: str, data: bytes, extract_tables: bool = True) -> List[Block]:
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext == "pdf":
        return _read_pdf(data, extract_tables)
    if ext == "docx":
        return _read_docx(data)
    if ext in ("txt", "md"):
        return _read_text(data)
    raise ValueError(f"unsupported file type '.{ext}'")


# --------------------------------------------------------------------------
# Chunking
# --------------------------------------------------------------------------
def _split_long(paragraph: str, size: int) -> List[str]:
    """Break an over-long paragraph on sentence/line boundaries."""
    if len(paragraph) <= size:
        return [paragraph]
    pieces, cur = [], ""
    for sent in re.split(r"(?<=[.!?])\s+|\n", paragraph):
        sent = sent.strip()
        if not sent:
            continue
        while len(sent) > size:  # a single monster sentence: hard split
            if cur:
                pieces.append(cur)
                cur = ""
            pieces.append(sent[:size])
            sent = sent[size:]
        if cur and len(cur) + len(sent) + 1 > size:
            pieces.append(cur)
            cur = sent
        else:
            cur = f"{cur} {sent}".strip()
    if cur:
        pieces.append(cur)
    return pieces


def split_text(text: str, size: int = 1000, overlap: int = 150) -> List[str]:
    text = text.replace("\r\n", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: List[str] = []
    cur = ""
    for para in paragraphs:
        for piece in _split_long(para, size):
            if cur and len(cur) + len(piece) + 2 > size:
                chunks.append(cur)
                tail = cur[-overlap:] if overlap else ""
                if " " in tail:
                    tail = tail[tail.find(" ") + 1:]  # don't start mid-word
                if tail and len(tail) + len(piece) + 2 <= size:
                    cur = f"{tail}\n\n{piece}"
                else:
                    cur = piece
            else:
                cur = f"{cur}\n\n{piece}" if cur else piece
    if cur:
        chunks.append(cur)
    return chunks


def split_table(markdown: str, size: int = 1000) -> List[str]:
    """Keep tables whole when possible; otherwise split by rows and repeat the header."""
    if len(markdown) <= size:
        return [markdown]
    lines = markdown.split("\n")
    header = "\n".join(lines[:2]) if len(lines) > 2 else ""
    rows = lines[2:] if header else lines
    chunks, cur = [], header
    for row in rows:
        row = row[:size]
        if cur and len(cur) + len(row) + 1 > size and cur != header:
            chunks.append(cur)
            cur = header
        cur = f"{cur}\n{row}" if cur else row
    if cur and cur != header:
        chunks.append(cur)
    return chunks


# --------------------------------------------------------------------------
# Public entry point
# --------------------------------------------------------------------------
def build_chunks(
    files: List[Tuple[str, bytes]],
    extract_tables: bool = True,
    chunk_size: int = 1000,
    overlap: int = 150,
    on_file: Optional[Callable[[str], None]] = None,
) -> Tuple[List[Chunk], List[str]]:
    """Return (chunks, warnings). Never raises for a bad file; it just warns."""
    chunks: List[Chunk] = []
    warnings: List[str] = []
    seen = set()

    for name, data in files:
        if on_file:
            on_file(name)
        try:
            blocks = read_file(name, data, extract_tables)
        except Exception as exc:
            warnings.append(f"Couldn't read **{name}**: {exc}")
            continue

        text_chars = sum(len(t) for _, kind, t in blocks if kind == "text")
        if text_chars < 30:
            warnings.append(
                f"**{name}** has no selectable text (probably a scanned PDF). "
                "OCR isn't included to keep this app light, so it was skipped.")
            continue

        for page, kind, text in blocks:
            pieces = (split_table(text, chunk_size) if kind == "table"
                      else split_text(text, chunk_size, overlap))
            for piece in pieces:
                piece = piece.strip()
                if len(piece) < 20:
                    continue
                key = hashlib.md5(" ".join(piece.lower().split()).encode()).hexdigest()
                if key in seen:
                    continue
                seen.add(key)
                chunks.append(Chunk(piece, name, page, kind))
    return chunks, warnings
