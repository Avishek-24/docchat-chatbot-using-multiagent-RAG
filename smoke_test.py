"""Offline self-check: no Ollama, no internet, no model download needed.

    python smoke_test.py
"""
import io

from docchat.agents import OUT_OF_SCOPE, VerificationAgent, run_pipeline
from docchat.loader import build_chunks, split_table, split_text
from docchat.retriever import HybridRetriever


class FakeLLM:
    """Stands in for Ollama so the whole pipeline can be tested offline."""

    def __init__(self, relevance="CAN_ANSWER", verdicts=None):
        self.relevance = relevance
        self.verdicts = list(verdicts or [])
        self.stream_calls = 0

    def complete(self, prompt, system=None, temperature=0.0, max_tokens=400, json_mode=False):
        if json_mode:
            return self.verdicts.pop(0) if self.verdicts else (
                '{"supported": true, "unsupported_claims": [], "contradictions": [], '
                '"relevant": true, "notes": "ok"}')
        return self.relevance

    def stream(self, prompt, system=None, temperature=0.2, max_tokens=700):
        self.stream_calls += 1
        yield from ["The PUE ", "in 2022 ", f"was 1.21 [1] (draft {self.stream_calls})."]


def make_files():
    import docx
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Data center efficiency report. Singapore 2nd facility PUE in 2022 was 1.21.\n"
                               "Asia Pacific carbon-free energy in 2023 was 12 percent.")
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "Waste diversion rate improved. Water consumption rose in 2023.")
    pdf_bytes = pdf.tobytes()
    blank = pymupdf.open()
    blank.new_page()
    blank_pdf = blank.tobytes()

    d = docx.Document()
    d.add_paragraph("Quarterly retention memo. Repeat customers spend more than new customers.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Segment", "Revenue"
    t.cell(1, 0).text, t.cell(1, 1).text = "Loyal", "1200"
    buf = io.BytesIO()
    d.save(buf)

    return [
        ("report.pdf", pdf_bytes),
        ("memo.docx", buf.getvalue()),
        ("notes.txt", "Plain text notes about onboarding and KYC checks.\n\nSecond paragraph.".encode()),
        ("readme.md", b"# Title\n\nSome markdown body text for testing the loader."),
        ("broken.pdf", b"this is not really a pdf"),
        ("empty.txt", b"   "),
        ("scan.pdf", blank_pdf),
    ]


def main():
    # 1. chunking edge cases
    assert split_text("short text that is long enough to keep") != []
    assert all(len(c) <= 1000 for c in split_text("word " * 2000, 1000, 150))
    assert all(len(c) <= 1000 for c in split_text("x" * 5000, 1000, 150))
    table = "| a | b |\n| --- | --- |\n" + "\n".join(f"| {i} | {i*2} |" for i in range(400))
    parts = split_table(table, 500)
    assert len(parts) > 1 and all(p.startswith("| a | b |") for p in parts)

    # 2. loading: good files, junk files, duplicates. Must never raise.
    files = make_files() + [("report.pdf", make_files()[0][1])]  # same file twice
    chunks, warnings = build_chunks(files)
    names = {c.source for c in chunks}
    assert {"report.pdf", "memo.docx", "notes.txt", "readme.md"} <= names, names
    assert any("broken.pdf" in w for w in warnings), warnings
    assert any("empty.txt" in w for w in warnings), warnings
    assert any("scan.pdf" in w and "selectable text" in w for w in warnings), warnings
    assert any(c.kind == "table" and "Loyal" in c.text for c in chunks)
    texts = [" ".join(c.text.split()) for c in chunks]
    assert len(texts) == len(set(texts)), "duplicate chunks were not removed"
    print(f"loader ok: {len(chunks)} chunks, {len(warnings)} warnings")

    # 3. retrieval (keyword-only mode, as when the embedder is unavailable)
    retriever = HybridRetriever(chunks, embedder=None)
    hits = retriever.search("What was the Singapore PUE in 2022?", k=3)
    assert hits and "1.21" in hits[0].chunk.text, [h.chunk.text[:60] for h in hits]
    assert retriever.search("???", k=3) == []
    print("retriever ok:", retriever.mode)

    # 3b. regression: a tiny corpus (one short file) must still return hits
    tiny, _ = build_chunks([("one.txt", b"Singapore PUE in 2022 was 1.21. Short note only.")])
    tiny_hits = HybridRetriever(tiny, None).search("What was the PUE in 2022?", k=3)
    assert tiny_hits and "1.21" in tiny_hits[0].chunk.text, "BM25 returned nothing on a tiny corpus"
    two, _ = build_chunks([("a.txt", b"Alpha document about apples and pears and fruit."),
                           ("b.txt", b"Beta document about bolts and nuts and hardware.")])
    assert HybridRetriever(two, None).search("apples", k=2)[0].chunk.source == "a.txt"

    # 4. pipeline: in-scope question, verified first time
    events = list(run_pipeline("PUE in 2022?", retriever, FakeLLM()))
    done = events[-1].payload
    assert events[-1].kind == "done" and "1.21" in done.answer
    assert done.verification["verdict"] == "verified" and done.revisions == 0
    assert any(e.kind == "token" for e in events)

    # 5. pipeline: out of scope question is refused
    done = list(run_pipeline("Who won the cricket world cup?", retriever, FakeLLM(relevance="NO_MATCH")))[-1].payload
    assert done.answer == OUT_OF_SCOPE and done.hits == []

    # 6. pipeline: self-correction loop actually retries, then stops
    bad = '{"supported": false, "unsupported_claims": ["1.21 is wrong"], "contradictions": [], "relevant": true, "notes": "x"}'
    llm = FakeLLM(verdicts=[bad, bad, bad])
    events = list(run_pipeline("PUE in 2022?", retriever, llm, max_retries=2))
    done = events[-1].payload
    assert done.revisions == 2 and llm.stream_calls == 3 and done.verification["verdict"] == "issues"
    assert sum(e.kind == "reset" for e in events) == 2
    llm = FakeLLM(verdicts=[bad])  # second check passes
    done = list(run_pipeline("PUE in 2022?", retriever, llm, max_retries=2))[-1].payload
    assert done.revisions == 1 and done.verification["verdict"] == "verified"

    # 7. verifier fails closed on garbage and tolerates fenced JSON
    v = VerificationAgent(FakeLLM())
    assert v.parse("lol not json")["verdict"] == "unverified"
    assert v.parse("")["verdict"] == "unverified"
    assert v.parse('```json\n{"supported": "YES", "unsupported_claims": [], "contradictions": []}\n```')["verdict"] == "verified"
    assert v.parse('{"supported": false}')["verdict"] == "unverified"  # no reasons given -> no pointless retry

    # 8. verification can be switched off
    done = list(run_pipeline("PUE in 2022?", retriever, FakeLLM(), verify=False))[-1].payload
    assert done.verification is None

    print("\nAll checks passed. You're good to run:  streamlit run app.py")


if __name__ == "__main__":
    main()
